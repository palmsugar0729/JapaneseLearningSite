# -*- coding: utf-8 -*-
"""把单词表图片 OCR 出来的词条整理成 Excel。

数据来源：codes/vocab_data/<单元名>.json，每个文件是一个数组，每行 5 个元素：
    [日语, 读音, 声调, 词性, 中文意思]
缺失的字段填空字符串。读音留空 = 原词只写了假名，导出时自动填成单词本身。

输出：docs/vocabulary_files/<单元名>.xlsx

用法：
    python codes/build_vocab_excel.py              # 生成全部单元
    python codes/build_vocab_excel.py "UNIT 1"     # 只生成指定单元
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(ROOT, "codes", "vocab_data")
OUT_DIR = os.path.join(ROOT, "docs", "vocabulary_files")

HEADERS = ["日语", "读音", "声调", "词性", "中文意思"]

HEAD_FONT = Font(bold=True, size=11)
HEAD_FILL = PatternFill("solid", fgColor="D9EAD3")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WIDTHS = [22, 26, 10, 14, 34]


def load_unit(unit: str):
    """读取 codes/vocab_data/<unit>.json，校验每行都是 5 列。"""
    path = os.path.join(DATA_DIR, f"{unit}.json")
    with open(path, encoding="utf-8") as f:
        rows = json.load(f)
    for i, r in enumerate(rows, 1):
        if not isinstance(r, list) or len(r) != 5:
            raise ValueError(f"{unit}.json 第 {i} 行不是 5 列：{r!r}")
    return rows


def build(unit: str, rows) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = unit

    ws.append(HEADERS)
    for c in ws[1]:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER

    for word, reading, tone, pos, meaning in rows:
        # 原词只写了假名时没有读音，按约定填成单词本身
        ws.append([word, reading or word, tone, pos, meaning])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(HEADERS)):
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(vertical="center", wrap_text=False)

    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"

    os.makedirs(OUT_DIR, exist_ok=True)
    path = os.path.join(OUT_DIR, f"{unit}.xlsx")
    wb.save(path)
    return path


def list_units():
    if not os.path.isdir(DATA_DIR):
        return []
    return sorted(f[:-5] for f in os.listdir(DATA_DIR) if f.endswith(".json"))


if __name__ == "__main__":
    targets = sys.argv[1:] or list_units()
    if not targets:
        sys.exit(f"没有找到数据文件，请先把 OCR 结果写到 {DATA_DIR}")
    for unit in targets:
        rows = load_unit(unit)
        path = build(unit, rows)
        print(f"{path}  ({len(rows)} 条)")
